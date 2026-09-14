"""Generate the synthetic DOCX and PDF example documents.

The binary files are not committed; run `python examples/generate_examples.py` (or call
`ensure_examples()`), which the Dockerfile, CI and the test suite already do.
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

HERE = Path(__file__).parent

# Demo document: exactly three deadlines, one needing a receipt date.
REGULATORY_NOTICE = [
    "SYNTHETIC EXAMPLE — REGULATORY NOTICE (fictional, for demo purposes only)",
    "To: Northwind Payments FZE. From: Fictional Financial Services Authority.",
    "Re: Periodic compliance review 2026/Q3.",
    "1. Documents must be submitted within 14 days from receipt of this notice.",
    "2. The annual compliance certificate must be filed no later than 30 September 2026.",
    "3. Should you wish to request an extension, you must give at least 5 business days' "
    "prior written notice to the Authority.",
    "4. This notice does not affect any other obligation under the applicable regulations.",
]

# Lease: several deadlines including months and a 'before expiry' clause.
LEASE = [
    "SYNTHETIC EXAMPLE — COMMERCIAL LEASE EXTRACT (fictional, for demo purposes only)",
    "Between Fictional Properties Ltd (Landlord) and Sample Retail LLC (Tenant).",
    "1. Rent. The Tenant shall pay the rent within 7 days of the invoice date.",
    "2. Deposit. The security deposit shall be paid within 10 business days of signing.",
    "3. Break clause. The Tenant may terminate this Lease by giving not less than three (3) "
    "months' written notice.",
    "4. Renewal. Notice of intention to renew must be given at least 90 days prior to the expiry "
    "of the term.",
    "5. Handover. The premises must be vacated no later than 31 December 2027.",
    "6. Insurance. Evidence of insurance must be provided within 30 days following the "
    "commencement date.",
]

# Document that mentions dates but contains no obligations with deadlines.
NO_DEADLINES = [
    "SYNTHETIC EXAMPLE — MEETING MINUTES (fictional, for demo purposes only)",
    "The parties met on 3 March 2026 to discuss the roadmap.",
    "Attendees: A. Sample, B. Example. The meeting was chaired by A. Sample.",
    "It was agreed that the general approach is sound and that no further action is required.",
]


def write_docx(paragraphs: list[str], path: Path) -> None:
    doc = Document()
    doc.add_heading(paragraphs[0], level=1)
    for p in paragraphs[1:]:
        doc.add_paragraph(p)
    doc.save(path)


def write_pdf(paragraphs: list[str], path: Path) -> None:
    styles = getSampleStyleSheet()
    doc = SimpleDocTemplate(str(path), pagesize=A4, title=paragraphs[0])
    story = [Paragraph(paragraphs[0], styles["Heading2"]), Spacer(1, 12)]
    for p in paragraphs[1:]:
        story += [Paragraph(p, styles["BodyText"]), Spacer(1, 8)]
    doc.build(story)


def write_txt(paragraphs: list[str], path: Path) -> None:
    path.write_text("\n\n".join(paragraphs) + "\n", encoding="utf-8")


GENERATED = {
    "regulatory_notice.pdf": (write_pdf, REGULATORY_NOTICE),
    "commercial_lease.docx": (write_docx, LEASE),
    "meeting_minutes_no_deadlines.docx": (write_docx, NO_DEADLINES),
}


def ensure_examples(force: bool = False) -> list[Path]:
    """Create any missing example file; returns the paths that were written."""
    written: list[Path] = []
    for name, (writer, paragraphs) in GENERATED.items():
        path = HERE / name
        if force or not path.exists():
            writer(paragraphs, path)
            written.append(path)
    return written


if __name__ == "__main__":
    for path in ensure_examples(force=True):
        print("wrote", path.relative_to(HERE.parent))
