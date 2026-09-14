"""Turn an uploaded file (PDF / DOCX / plain text) into plain text."""

from __future__ import annotations

import io
import re
import zipfile

from docx import Document
from docx.opc.exceptions import PackageNotFoundError
from pypdf import PdfReader
from pypdf.errors import PyPdfError

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md"}
MAX_FILE_BYTES = 10 * 1024 * 1024


class UnsupportedFileError(ValueError):
    pass


class EmptyDocumentError(ValueError):
    pass


def _normalize(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\xa0", " ")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def parse_pdf(data: bytes) -> str:
    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            raise UnsupportedFileError("PDF is encrypted")
        pages = [page.extract_text() or "" for page in reader.pages]
    except PyPdfError as exc:  # malformed file
        raise UnsupportedFileError(f"Cannot read PDF: {exc}") from exc
    text = _normalize("\n".join(pages))
    if not text:
        raise EmptyDocumentError(
            "PDF contains no extractable text (scanned image? OCR is not supported)"
        )
    return text


def parse_docx(data: bytes) -> str:
    try:
        doc = Document(io.BytesIO(data))
    except (PackageNotFoundError, zipfile.BadZipFile, KeyError, ValueError) as exc:
        raise UnsupportedFileError(f"Cannot read DOCX: {exc}") from exc
    parts: list[str] = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    text = _normalize("\n".join(parts))
    if not text:
        raise EmptyDocumentError("DOCX contains no text")
    return text


def parse_text(data: bytes) -> str:
    for enc in ("utf-8", "cp1251", "latin-1"):
        try:
            text = data.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    else:  # pragma: no cover - latin-1 never fails
        text = data.decode("latin-1", errors="replace")
    text = _normalize(text)
    if not text:
        raise EmptyDocumentError("Text is empty")
    return text


def parse_upload(filename: str, data: bytes) -> str:
    if len(data) > MAX_FILE_BYTES:
        raise UnsupportedFileError("File is larger than 10 MB")
    name = (filename or "").lower()
    ext = name[name.rfind(".") :] if "." in name else ""
    if ext not in SUPPORTED_EXTENSIONS:
        raise UnsupportedFileError(
            f"Unsupported file type '{ext or 'none'}'. Supported: PDF, DOCX, TXT, MD"
        )
    if ext == ".pdf":
        return parse_pdf(data)
    if ext == ".docx":
        return parse_docx(data)
    return parse_text(data)
